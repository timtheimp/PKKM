import json
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))
from export_workbook import _reject_source_overwrite, _excel_safe_value, _write_table, export_workbook


class WorkbookExportTests(unittest.TestCase):
    def test_excel_formula_values_are_neutralized(self):
        for prefix in ("=", "+", "-", "@"):
            value = prefix + "HYPERLINK(\"https://example.invalid\")"
            self.assertEqual(_excel_safe_value(value), "'" + value)
        self.assertEqual(_excel_safe_value("ordinary text"), "ordinary text")
        workbook = Workbook()
        try:
            sheet = workbook.active
            _write_table(sheet, ["=FIELD"], [["safe"]])
            self.assertEqual(sheet["A1"].value, "'=FIELD")
        finally:
            workbook.close()

    def test_export_rejects_source_overwrite(self):
        source = Path("kingdom.json")
        with self.assertRaisesRegex(ValueError, "Refusing to overwrite source JSON"):
            _reject_source_overwrite(source, source)

    def test_export_preserves_core_sections(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "kingdom.json"
            destination = root / "export.xlsx"
            source.write_text(json.dumps({
                "schema_version": 10,
                "name": "Exported Kingdom",
                "treasury_bp": 12,
                "settlements": [{
                    "id": "settlement.one",
                    "name": "One",
                    "population": 10,
                    "districts": 1,
                    "building_inventory": [{"name": "Academy", "count": 1}],
                    "map": {"placements": [{"building_name": "Academy", "row": 0, "column": 0, "width": 1, "height": 1}]},
                }],
                "greenbelt_map": {"hexes": [{"q": 1, "r": 2, "terrain": "Plains, River", "claimed": True, "improvements": ["farm"]}]},
                "kingdom_events": [{"id": "event-1", "roll": 10}],
                "valuable_item_sales": [{"item_name": "Recovered Crown", "bp_received": 8}],
            }), encoding="utf-8")
            export_workbook(source, destination)
            workbook = load_workbook(destination, read_only=True, data_only=True)
            try:
                self.assertEqual(set(workbook.sheetnames), {"Kingdom", "Rules", "Settlements", "Buildings", "Placements", "Greenbelt", "Events", "Construction", "Valuable Sales", "Calendar", "Turn Actions"})
                self.assertEqual(workbook["Kingdom"]["B2"].value, 10)
                self.assertEqual(workbook["Settlements"]["B2"].value, "One")
                self.assertEqual(workbook["Buildings"]["B2"].value, "Academy")
                self.assertEqual(workbook["Greenbelt"]["C2"].value, "Plains, River")
                self.assertEqual(workbook["Events"]["B2"].value, 10)
                self.assertIn("Recovered Crown", [cell.value for row in workbook["Valuable Sales"].iter_rows() for cell in row])
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
