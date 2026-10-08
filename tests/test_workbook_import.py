import json
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))
from import_workbook import import_workbook


class WorkbookImportTests(unittest.TestCase):
    def make_workbook(self, path: Path) -> None:
        workbook = Workbook()
        kingdom = workbook.active
        kingdom.title = "Kingdom"
        kingdom["A1"] = "Imported Kingdom"
        kingdom["C2"] = "CE"
        kingdom["E2"] = "=B54"
        kingdom["B10"] = "Standard"
        kingdom["B11"] = "Yes"
        kingdom["B14"] = "Light"
        kingdom["B15"] = "No"
        kingdom["B18"] = 6
        kingdom["B23"] = 3
        kingdom["D29"] = "Capital"
        kingdom["J38"] = "=SUM(J29:J37)"
        kingdom["E11"] = "Baron Name"
        kingdom["G11"] = 2
        kingdom["H11"] = "Economy"
        kingdom["E17"] = "Magister Name"
        kingdom["G17"] = 3
        kingdom["D54"] = 4
        kingdom["E54"] = 5
        kingdom["F54"] = 6

        capital = workbook.create_sheet("Capital")
        capital["A1"] = "Capital"
        capital["D3"] = 2
        capital["D4"] = 500
        capital["A11"] = "Castle"
        capital["B11"] = 1
        capital["E37"] = "=1"

        capital_map = workbook.create_sheet("Capital Map")
        capital_map.merge_cells("B1:G1")
        capital_map["B1"] = "Hill"
        capital_map.merge_cells("F6:G7")
        capital_map["F6"] = "Castle"

        calendar = workbook.create_sheet("Calendar")
        calendar["B1"] = "Month"
        calendar["C1"] = "Holiday"
        calendar["B2"] = "Pharast (March)"
        calendar["C2"] = "Founding Day"
        calendar["A14"] = "1 yr!"
        calendar["B14"] = "Pharast (March)"

        workbook.create_sheet("Turn")
        data = workbook.create_sheet("Data")
        data["A1"] = "Improvements"
        workbook.defined_names.add(DefinedName("IMP", attr_text="Data!$A$2:$T$2"))
        workbook.save(path)

    def test_import_maps_inputs_and_reports_unsupported_formulas(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.xlsx"
            self.make_workbook(source)

            kingdom, report = import_workbook(source, catalog_names={"Castle"})

        self.assertEqual(kingdom["schema_version"], 9)
        self.assertEqual(kingdom["name"], "Imported Kingdom")
        self.assertEqual(kingdom["rules"]["alignment"], "CE")
        self.assertEqual(kingdom["rules"]["laws"]["promotion"], 2)
        self.assertEqual(kingdom["rules"]["laws"]["holiday_interval_months"], 6)
        self.assertEqual(kingdom["unrest"], 3)
        self.assertEqual(kingdom["settlements"][0]["building_inventory"], [{"name": "Castle", "count": 1}])
        self.assertEqual(kingdom["settlements"][0]["map"]["placements"], [
            {"building_name": "Castle", "row": 5, "column": 5, "width": 2, "height": 2}
        ])
        self.assertEqual(kingdom["calendar_notes"], [{
            "entry_id": "calendar.row.2",
            "holiday": "Founding Day",
            "kingdom_upgrades": "",
            "events": "",
            "other": "",
        }])
        self.assertIn(
            {"sheet": "Kingdom", "cell": "E2"},
            [{"sheet": item["sheet"], "cell": item["cell"]} for item in report["unsupported"]],
        )
        self.assertGreater(report["summary"]["mapped_cells"], 0)
        self.assertGreaterEqual(report["summary"]["unsupported_formulas"], 2)
        self.assertEqual(report["sheets"]["Turn"]["status"], "unsupported_manual_workflow")


if __name__ == "__main__":
    unittest.main()
