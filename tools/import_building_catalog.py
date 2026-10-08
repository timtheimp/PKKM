#!/usr/bin/env python3
"""Read the workbook's IMP named range into an auditable JSON catalog."""

from __future__ import annotations

import argparse
import json
import math
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

EXPECTED_HEADERS = [
    "Improvements", "Cost", "Lots", "Economy", "Loyalty", "Stability", "Defense", "Unrest",
    "Base Value", "Discount", "Magic Item", "Upgrade From", "Upgrade To", "Corruption", "Crime",
    "Law", "Lore", "Society", "Productivity", "Fame",
]
KEYS = [
    "name", "cost", "lots", "economy", "loyalty", "stability", "defense", "unrest", "base_value",
    "discounts", "magic_item", "upgrade_from", "upgrade_to", "corruption", "crime", "law", "lore",
    "society", "productivity", "fame",
]
NUMERIC_KEYS = {
    "cost", "lots", "economy", "loyalty", "stability", "defense", "unrest", "base_value",
    "corruption", "crime", "law", "lore", "society", "productivity", "fame",
}
TEXT_KEYS = {"discounts", "magic_item", "upgrade_from", "upgrade_to"}
NON_BUILDING_KINDS = {"basic", "town", "city", "district"}


def json_cell(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        if isinstance(value, float) and not math.isfinite(value):
            raise ValueError(f"Workbook contains a non-finite number: {value!r}")
        if float(value).is_integer():
            return int(value)
        return value
    if isinstance(value, str):
        return value
    raise TypeError(f"Unsupported workbook cell value: {type(value).__name__}")


def export_catalog(workbook_path: Path) -> dict[str, Any]:
    workbook = load_workbook(workbook_path, data_only=False, read_only=False)
    try:
        if "IMP" not in workbook.defined_names:
            raise ValueError("Workbook is missing the IMP named range")
        reference = workbook.defined_names["IMP"].attr_text
        sheet_name, cell_range = reference.rsplit("!", 1)
        sheet_name = sheet_name.strip("'")
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"IMP refers to missing sheet {sheet_name!r}")
        min_col, min_row, max_col, max_row = range_boundaries(cell_range)
        if max_col - min_col + 1 != len(EXPECTED_HEADERS) or min_row < 2:
            raise ValueError(f"Unexpected IMP range: {reference}")

        sheet = workbook[sheet_name]
        headers = [sheet.cell(min_row - 1, column).value for column in range(min_col, max_col + 1)]
        if headers != EXPECTED_HEADERS:
            raise ValueError(f"Unexpected IMP headers: {headers!r}")

        records = []
        for source_row in range(min_row, max_row + 1):
            raw = [json_cell(sheet.cell(source_row, column).value) for column in range(min_col, max_col + 1)]
            if all(value is None for value in raw):
                continue
            label = raw[0]
            if not isinstance(label, str) or not label.strip():
                kind = "unclassified"
                name = "" if label is None else str(label)
            else:
                name = label.strip()
                marker = name.casefold()
                kind = marker if marker in NON_BUILDING_KINDS else "improvement"

            record: dict[str, Any] = {
                "source_row": source_row,
                "kind": kind,
                "name": name,
                "source_values": raw,
            }
            if kind == "improvement":
                mapped = dict(zip(KEYS, raw, strict=True))
                if (not isinstance(mapped["cost"], int) or isinstance(mapped["cost"], bool)
                        or not isinstance(mapped["lots"], int) or isinstance(mapped["lots"], bool)):
                    raise ValueError(f"Improvement row {source_row} must have integer Cost and Lots")
                building: dict[str, Any] = {}
                for key in NUMERIC_KEYS:
                    value = mapped[key]
                    if value is None:
                        building[key] = 0
                    elif isinstance(value, int) and not isinstance(value, bool):
                        building[key] = value
                    else:
                        raise ValueError(f"Improvement row {source_row} has a non-integer {key}: {value!r}")
                for key in TEXT_KEYS:
                    value = mapped[key]
                    if value is None:
                        building[key] = ""
                    elif isinstance(value, str):
                        building[key] = value
                    else:
                        raise ValueError(f"Improvement row {source_row} has non-text {key}: {value!r}")
                record["building"] = building
            records.append(record)

        return {
            "schema_version": 1,
            "source_workbook": workbook_path.name,
            "source_sheet": sheet_name,
            "named_range": "IMP",
            "source_range": reference,
            "headers": headers,
            "records": records,
        }
    finally:
        workbook.close()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path, help="Source .xlsx workbook; read only")
    parser.add_argument("output", type=Path, help="Destination JSON asset")
    args = parser.parse_args()
    catalog = export_catalog(args.workbook)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(catalog, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Wrote {len(catalog['records'])} source rows from {catalog['source_range']} to {args.output}")


if __name__ == "__main__":
    main()
