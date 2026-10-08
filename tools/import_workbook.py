#!/usr/bin/env python3
"""Import supported workbook inputs into a schema-v9 kingdom JSON and report gaps."""

from __future__ import annotations

import argparse
import json
import math
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

PROMOTION_VALUES = {"None": 0, "Token": 1, "Standard": 2, "Aggressive": 3, "Expansionist": 4}
TAXATION_VALUES = {"None": 0, "Light": 1, "Normal": 2, "Heavy": 3, "Overwhelming": 4}
LEADERSHIP_ROWS = {
    "councilor": 13, "general": 14, "grand_diplomat": 15, "high_priest": 16,
    "magister": 17, "marshal": 18, "royal_enforcer": 19, "spymaster": 20,
    "treasurer": 21, "warden": 22,
}
STAT_KEYS = ("economy", "loyalty", "stability")


def _value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        if isinstance(value, float) and not math.isfinite(value):
            raise ValueError(f"Workbook contains a non-finite number: {value!r}")
        return int(value) if float(value).is_integer() else value
    if isinstance(value, str):
        return value
    return str(value)


def _formula(value: Any) -> bool:
    return isinstance(value, str) and value.startswith("=")


def _text(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(str(value).replace("\r", "\n").split())


def _blank_rules() -> dict[str, Any]:
    slot = {"filled": False, "bonus": 0}
    return {
        "alignment": "N",
        "territory": {
            "kingdom_size": 0, "total_city_districts": 0, "control_dc_other": 0,
            "city_count": 0, "total_city_population": 0,
        },
        "laws": {
            "promotion": 0, "has_cathedral": False, "taxation": 0,
            "has_waterfront": False, "holiday_interval_months": 0,
        },
        "leadership": {
            "baron_bonus": 0,
            "baron_benefits": {"economy": True, "loyalty": False, "stability": False},
            "spymaster_benefit": 0,
            **{name: dict(slot) for name in LEADERSHIP_ROWS},
        },
        "stats": {
            "economy": {"events": 0, "improvements": 0, "other": 0},
            "loyalty": {"events": 0, "improvements": 0, "other": 0},
            "stability": {"events": 0, "improvements": 0, "other": 0},
        },
    }


def _empty_kingdom() -> dict[str, Any]:
    return {
        "schema_version": 9,
        "name": "Imported Kingdom",
        "turn": 1,
        "treasury_bp": 0,
        "unrest": 0,
        "rules": _blank_rules(),
        "settlements": [],
        "buildings": [],
        "turn_progress": [],
        "calendar_notes": [],
    }


def _focus(value: Any) -> int:
    text = _text(value).casefold()
    return {"economy": 0, "loyalty": 1, "stability": 2}.get(text, 0)


def _coverage(value: Any) -> dict[str, bool]:
    text = _text(value).casefold()
    return {
        "economy": "economy" in text or text in {"e/l", "e/s", "e/l/s"},
        "loyalty": "loyalty" in text or text in {"e/l", "l/s", "e/l/s"},
        "stability": "stability" in text or text in {"e/s", "l/s", "e/l/s"},
    }


def import_workbook(workbook_path: Path, catalog_names: set[str] | None = None) -> tuple[dict[str, Any], dict[str, Any]]:
    workbook_path = Path(workbook_path)
    workbook = load_workbook(workbook_path, data_only=False, read_only=False)
    kingdom = _empty_kingdom()
    mapped: list[dict[str, Any]] = []
    unsupported: list[dict[str, Any]] = []
    sheets: dict[str, dict[str, Any]] = {}
    catalog_names = catalog_names or set()

    def mapped_cell(sheet: str, cell: str, field: str, value: Any) -> None:
        mapped.append({"sheet": sheet, "cell": cell, "field": field, "value": _value(value)})

    def sheet_cell(sheet: str, cell: str, field: str) -> Any:
        value = workbook[sheet][cell].value
        if _formula(value):
            return None
        mapped_cell(sheet, cell, field, value)
        return _value(value)

    try:
        for sheet in workbook.worksheets:
            formula_cells = []
            for row in sheet.iter_rows():
                for cell in row:
                    if _formula(cell.value):
                        formula_cells.append(cell.coordinate)
                        unsupported.append({
                            "sheet": sheet.title,
                            "cell": cell.coordinate,
                            "formula": cell.value,
                            "reason": "Derived formula is not evaluated during import",
                        })
            status = "supported_inputs"
            if sheet.title == "Turn":
                status = "unsupported_manual_workflow"
            elif sheet.title == "Data":
                status = "reference_catalog_handled_separately"
            elif sheet.title not in {"Kingdom", "Calendar", "Capital", "City #", "Capital Map", "City # Map", "Example Map", "Multi-District Map"}:
                status = "unsupported_sheet"
            sheets[sheet.title] = {
                "status": status,
                "rows": sheet.max_row,
                "columns": sheet.max_column,
                "merged_ranges": len(sheet.merged_cells.ranges),
                "formula_cells": len(formula_cells),
            }

        kingdom_sheet = workbook["Kingdom"]
        name = sheet_cell("Kingdom", "A1", "kingdom.name")
        if isinstance(name, str) and name.strip():
            kingdom["name"] = name.strip()
        alignment = sheet_cell("Kingdom", "C2", "rules.alignment")
        if isinstance(alignment, str) and alignment.strip():
            kingdom["rules"]["alignment"] = alignment.strip()

        direct_numeric = {
            ("E2", "rules.territory.kingdom_size"),
            ("H2", "rules.territory.control_dc_other"),
            ("B23", "unrest"),
        }
        for cell, field in direct_numeric:
            value = sheet_cell("Kingdom", cell, field)
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                target, key = field.rsplit(".", 1) if "." in field else ("", field)
                if field == "unrest":
                    kingdom["unrest"] = int(value)
                else:
                    parent = kingdom["rules"]
                    for part in field.removeprefix("rules.").split(".")[:-1]:
                        parent = parent[part]
                    parent[key] = int(value)

        law_values = {
            "B10": ("rules.laws.promotion", PROMOTION_VALUES),
            "B14": ("rules.laws.taxation", TAXATION_VALUES),
        }
        for cell, (field, choices) in law_values.items():
            value = sheet_cell("Kingdom", cell, field)
            if isinstance(value, str) and value in choices:
                kingdom["rules"]["laws"][field.rsplit(".", 1)[1]] = choices[value]
        cathedral = sheet_cell("Kingdom", "B11", "rules.laws.has_cathedral")
        if isinstance(cathedral, str):
            kingdom["rules"]["laws"]["has_cathedral"] = cathedral.casefold() == "yes"
        waterfront = sheet_cell("Kingdom", "B15", "rules.laws.has_waterfront")
        if isinstance(waterfront, str):
            kingdom["rules"]["laws"]["has_waterfront"] = waterfront.casefold() == "yes"
        holiday = sheet_cell("Kingdom", "B18", "rules.laws.holiday_interval_months")
        if isinstance(holiday, (int, float)) and not isinstance(holiday, bool):
            kingdom["rules"]["laws"]["holiday_interval_months"] = int(holiday)

        leadership = kingdom["rules"]["leadership"]
        baron_name = sheet_cell("Kingdom", "E11", "rules.leadership.baron.name")
        baron_bonus = sheet_cell("Kingdom", "G11", "rules.leadership.baron_bonus")
        baron_benefits = sheet_cell("Kingdom", "H11", "rules.leadership.baron_benefits")
        if isinstance(baron_bonus, (int, float)) and not isinstance(baron_bonus, bool):
            leadership["baron_bonus"] = int(baron_bonus)
        if baron_benefits is not None:
            leadership["baron_benefits"] = _coverage(baron_benefits)
        for role, row in LEADERSHIP_ROWS.items():
            name_value = sheet_cell("Kingdom", f"E{row}", f"rules.leadership.{role}.name")
            bonus_value = sheet_cell("Kingdom", f"G{row}", f"rules.leadership.{role}.bonus")
            leadership[role] = {
                "filled": isinstance(name_value, str) and bool(name_value.strip()),
                "bonus": int(bonus_value) if isinstance(bonus_value, (int, float)) and not isinstance(bonus_value, bool) else 0,
            }
        spymaster_benefit = sheet_cell("Kingdom", "H20", "rules.leadership.spymaster_benefit")
        if spymaster_benefit is not None:
            leadership["spymaster_benefit"] = _focus(spymaster_benefit)

        for stat, row in zip(STAT_KEYS, (5, 6, 7)):
            events = sheet_cell("Kingdom", f"C{row}", f"rules.stats.{stat}.events")
            other = sheet_cell("Kingdom", f"J{row}", f"rules.stats.{stat}.other")
            if isinstance(events, (int, float)) and not isinstance(events, bool):
                kingdom["rules"]["stats"][stat]["events"] = int(events)
            if isinstance(other, (int, float)) and not isinstance(other, bool):
                kingdom["rules"]["stats"][stat]["other"] = int(other)

        for row in range(29, 38):
            ref = sheet_cell("Kingdom", f"D{row}", f"settlements[{len(kingdom['settlements'])}].source_sheet")
            if not isinstance(ref, str) or not ref.strip() or ref not in workbook.sheetnames:
                continue
            source = workbook[ref]
            settlement: dict[str, Any] = {
                "name": _text(source["A1"].value) or ref,
                "population": 0,
                "districts": 0,
                "economy": 0,
                "loyalty": 0,
                "stability": 0,
                "defense": 0,
                "building_inventory": [],
                "apply_catalog_stat_effects": False,
                "map": {"rows": 0, "columns": 0, "placements": [], "labels": []},
            }
            mapped_cell(ref, "A1", f"settlements[{len(kingdom['settlements'])}].name", source["A1"].value)
            for cell, key in (("D3", "districts"), ("D4", "population")):
                value = source[cell].value
                if not _formula(value) and isinstance(value, (int, float)) and not isinstance(value, bool):
                    settlement[key] = int(value)
                    mapped_cell(ref, cell, f"settlements[{len(kingdom['settlements'])}].{key}", value)
            for inventory_row in range(11, 37):
                item = source[f"A{inventory_row}"].value
                count = source[f"B{inventory_row}"].value
                if isinstance(item, str) and item.strip() and isinstance(count, (int, float)) and not isinstance(count, bool) and int(count) > 0:
                    settlement["building_inventory"].append({"name": item.strip(), "count": int(count)})
                    mapped_cell(ref, f"A{inventory_row}", f"settlements[{len(kingdom['settlements'])}].building_inventory", item)
                    mapped_cell(ref, f"B{inventory_row}", f"settlements[{len(kingdom['settlements'])}].building_inventory", count)
            map_name = f"{ref} Map"
            if map_name in workbook.sheetnames:
                settlement["map"] = _import_map(workbook[map_name], catalog_names, mapped_cell, ref, len(kingdom["settlements"]))
            kingdom["settlements"].append(settlement)

        calendar = workbook["Calendar"]
        for row in range(2, calendar.max_row + 1):
            month = calendar[f"B{row}"].value
            if not isinstance(month, str) or not month.strip():
                continue
            marker = calendar[f"A{row}"].value
            mapped_cell("Calendar", f"B{row}", f"calendar_template.row_{row}.month", month)
            if marker is not None and not _formula(marker):
                mapped_cell("Calendar", f"A{row}", f"calendar_template.row_{row}.year_marker", marker)
            note = {"entry_id": f"calendar.row.{row}", "holiday": "", "kingdom_upgrades": "", "events": "", "other": ""}
            for column, key in (("C", "holiday"), ("D", "kingdom_upgrades"), ("E", "events"), ("F", "other")):
                value = calendar[f"{column}{row}"].value
                if _formula(value):
                    continue
                if value is not None:
                    note[key] = str(value)
                    mapped_cell("Calendar", f"{column}{row}", f"calendar_notes[{row}].{key}", value)
            if any(note[key] for key in ("holiday", "kingdom_upgrades", "events", "other")):
                kingdom["calendar_notes"].append(note)

        sheets["Calendar"]["calendar_entries"] = sum(
            1 for row in range(2, calendar.max_row + 1) if isinstance(calendar[f"B{row}"].value, str)
        )
        report = {
            "schema_version": 1,
            "source_workbook": workbook_path.name,
            "mapped": mapped,
            "unsupported": unsupported,
            "sheets": sheets,
            "summary": {
                "mapped_cells": len(mapped),
                "unsupported_formulas": len(unsupported),
                "settlements_imported": len(kingdom["settlements"]),
                "calendar_notes_imported": len(kingdom["calendar_notes"]),
            },
            "limitations": [
                "Formula cells are reported but not evaluated; imported values come from direct workbook inputs only.",
                "Turn is a manual procedure sheet and is reported without automated outcomes or checklist completion.",
                "Catalog reference data remains in the separate IMP-derived building_catalog.json asset.",
            ],
        }
        return kingdom, report
    finally:
        workbook.close()


def _import_map(sheet: Any, catalog_names: set[str], mapped_cell: Any, settlement_sheet: str, index: int) -> dict[str, Any]:
    result = {"rows": sheet.max_row, "columns": sheet.max_column, "placements": [], "labels": []}
    merged_anchors = {}
    for merged in sheet.merged_cells.ranges:
        min_col, min_row, max_col, max_row = range_boundaries(str(merged))
        merged_anchors[(min_row, min_col)] = (max_row - min_row + 1, max_col - min_col + 1)
    seen = set()
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is None or _formula(cell.value):
                continue
            min_col = cell.column
            min_row = cell.row
            dimensions = merged_anchors.get((min_row, min_col), (1, 1))
            if dimensions == (1, 1):
                covered_by_merged = any(
                    merged != (min_row, min_col)
                    and merged[0] <= min_row <= merged[0] + height - 1
                    and merged[1] <= min_col <= merged[1] + width - 1
                    for merged, (height, width) in merged_anchors.items()
                )
                if covered_by_merged:
                    continue
            key = (min_row, min_col)
            if key in seen:
                continue
            seen.add(key)
            text = _text(cell.value)
            if not text:
                continue
            height, width = dimensions
            payload = {"row": min_row - 1, "column": min_col - 1, "width": width, "height": height}
            if text in catalog_names:
                result["placements"].append({"building_name": text, **payload})
                mapped_cell(sheet.title, cell.coordinate, f"settlements[{index}].map.placements", cell.value)
            else:
                result["labels"].append({"text": text, **payload})
                mapped_cell(sheet.title, cell.coordinate, f"settlements[{index}].map.labels", cell.value)
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("output", type=Path, help="Destination schema-v9 kingdom JSON")
    parser.add_argument("--report", type=Path, required=True, help="Destination mapping report JSON")
    parser.add_argument("--catalog", type=Path, help="Optional building_catalog.json for map placement classification")
    args = parser.parse_args()
    catalog_names: set[str] = set()
    if args.catalog:
        catalog = json.loads(args.catalog.read_text(encoding="utf-8"))
        catalog_names = {record["name"] for record in catalog.get("records", []) if record.get("kind") == "improvement"}
    kingdom, report = import_workbook(args.workbook, catalog_names)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(kingdom, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"PKKM_WORKBOOK_IMPORT_PASS mapped={report['summary']['mapped_cells']} formulas={report['summary']['unsupported_formulas']} settlements={report['summary']['settlements_imported']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
