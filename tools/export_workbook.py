#!/usr/bin/env python3
"""Export schema-v10 kingdom JSON to a new, read-only-friendly XLSX report."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def _reject_source_overwrite(source: Path, destination: Path) -> None:
    if source.resolve() == destination.resolve():
        raise ValueError(f"Refusing to overwrite source JSON: {source}")


def _excel_safe_value(value: Any) -> Any:
    if isinstance(value, str) and value[:1] in {"=", "+", "-", "@"}:
        return "'" + value
    return value


def _write_table(ws: Any, headers: list[str], rows: list[list[Any]]) -> None:
    safe_headers = [_excel_safe_value(header) for header in headers]
    ws.append(safe_headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for row in rows:
        ws.append([_excel_safe_value(json.dumps(value, ensure_ascii=False)
            if isinstance(value, (dict, list)) else value) for value in row])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for index, header in enumerate(safe_headers, 1):
        width = max(len(str(header)), *(len(str(row[index - 1])) if index - 1 < len(row) and row[index - 1] is not None else 0 for row in rows), 10)
        ws.column_dimensions[get_column_letter(index)].width = min(width + 2, 42)


def _scalar_rows(data: dict[str, Any]) -> list[list[Any]]:
    rows: list[list[Any]] = []
    for key, value in data.items():
        if isinstance(value, (str, int, float, bool)) or value is None:
            rows.append([key, value])
    return rows


def export_workbook(source: Path, destination: Path) -> None:
    _reject_source_overwrite(source, destination)
    data = json.loads(source.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise ValueError("Kingdom JSON root must be an object")

    workbook = Workbook()
    kingdom = workbook.active
    kingdom.title = "Kingdom"
    _write_table(kingdom, ["Field", "Value"], _scalar_rows(data))

    rules = data.get("rules", {})
    if isinstance(rules, dict):
        sheet = workbook.create_sheet("Rules")
        rows: list[list[Any]] = []
        for section, values in rules.items():
            if isinstance(values, dict):
                for key, value in values.items():
                    if isinstance(value, dict):
                        for subkey, subvalue in value.items():
                            rows.append([section, key, subkey, subvalue])
                    else:
                        rows.append([section, key, "", value])
            else:
                rows.append(["", section, "", values])
        _write_table(sheet, ["Section", "Field", "Subfield", "Value"], rows)

    settlements = data.get("settlements", [])
    sheet = workbook.create_sheet("Settlements")
    settlement_rows: list[list[Any]] = []
    inventory_rows: list[list[Any]] = []
    placement_rows: list[list[Any]] = []
    for settlement in settlements if isinstance(settlements, list) else []:
        if not isinstance(settlement, dict):
            continue
        settlement_rows.append([settlement.get(k) for k in ("id", "name", "population", "districts", "economy", "loyalty", "stability", "defense", "apply_catalog_stat_effects")])
        for item in settlement.get("building_inventory", []) if isinstance(settlement.get("building_inventory"), list) else []:
            if isinstance(item, dict):
                inventory_rows.append([settlement.get("id"), item.get("name"), item.get("count")])
        mapping = settlement.get("map", {})
        for placement in mapping.get("placements", []) if isinstance(mapping, dict) else []:
            if isinstance(placement, dict):
                placement_rows.append([settlement.get("id"), placement.get("building_name"), placement.get("row"), placement.get("column"), placement.get("width"), placement.get("height")])
    _write_table(sheet, ["ID", "Name", "Population", "Districts", "Economy", "Loyalty", "Stability", "Defense", "Apply Catalog Stats"], settlement_rows)

    sheet = workbook.create_sheet("Buildings")
    _write_table(sheet, ["Settlement ID", "Building", "Count"], inventory_rows)
    sheet = workbook.create_sheet("Placements")
    _write_table(sheet, ["Settlement ID", "Building", "Row", "Column", "Width", "Height"], placement_rows)

    greenbelt = data.get("greenbelt_map", {})
    sheet = workbook.create_sheet("Greenbelt")
    hex_rows: list[list[Any]] = []
    if isinstance(greenbelt, dict):
        for item in greenbelt.get("hexes", []) if isinstance(greenbelt.get("hexes"), list) else []:
            if isinstance(item, dict):
                hex_rows.append([item.get(k) for k in ("q", "r", "terrain", "claimed", "claim_turn", "settlement_id", "terrain_notes", "improvements")])
    _write_table(sheet, ["Q", "R", "Terrain", "Claimed", "Claim Turn", "Settlement ID", "Terrain Notes", "Improvements"], hex_rows)

    for key, title in (("kingdom_events", "Events"), ("construction_transactions", "Construction"), ("valuable_item_sales", "Valuable Sales"), ("calendar_notes", "Calendar"), ("turn_actions", "Turn Actions")):
        values = data.get(key, [])
        sheet = workbook.create_sheet(title)
        if isinstance(values, list) and values and isinstance(values[0], dict):
            headers = sorted({field for value in values if isinstance(value, dict) for field in value})
            rows = [[value.get(field) for field in headers] for value in values if isinstance(value, dict)]
            _write_table(sheet, headers or ["Value"], rows or [[""]])
        else:
            _write_table(sheet, ["Value"], [[value] for value in values] if isinstance(values, list) else [[values]])

    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path, help="schema-v10 kingdom JSON")
    parser.add_argument("destination", type=Path, help="new XLSX report path")
    args = parser.parse_args()
    export_workbook(args.source, args.destination)
    print(json.dumps({"source": str(args.source), "destination": str(args.destination), "exported": True}))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
